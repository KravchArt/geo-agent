"""Fill an Excel workbook with answers returned by the GeoAgent API.

Each non-empty query is sent to ``POST /api/v1/chat`` as an independent
conversation. Successful answers are checkpointed back to the source workbook
after every row, so an interrupted run can be resumed without losing progress.
"""

from __future__ import annotations

import argparse
import asyncio
import os
import re
import uuid
from collections.abc import Sequence
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Protocol

import httpx
from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import ValidationError

from common.models import AgentResponse

DEFAULT_AGENT_URL = "http://localhost:8000/api/v1/chat"
DEFAULT_ANSWER_COLUMN = "answer"
DEFAULT_ANSWER_WITH_TOOLS_COLUMN = "answer_with_tools"
DEFAULT_API_CALL_LIMIT = 1
QUERY_COLUMN_CANDIDATES = ("user_query", "query", "request", "question", "prompt")
SUPPORTED_SUFFIXES = {".xlsx", ".xlsm"}


@dataclass(frozen=True, slots=True)
class PendingRow:
    excel_row: int
    query: str


@dataclass(frozen=True, slots=True)
class FillResult:
    attempted: int
    written: int
    skipped: int
    failed: int


@dataclass(frozen=True, slots=True)
class SuccessfulToolResult:
    tool_name: str
    content: str


@dataclass(frozen=True, slots=True)
class CollectedAnswer:
    answer: str
    successful_tool_results: tuple[SuccessfulToolResult, ...] = ()

    def with_tool_results(self) -> str:
        sections = [
            f"Tool result ({result.tool_name}):\n{result.content}"
            for result in self.successful_tool_results
        ]
        sections.append(f"Final answer:\n{self.answer}")
        return "\n\n".join(sections)


class AnswerClient(Protocol):
    async def answer(self, query: str, *, session_id: str) -> CollectedAnswer:
        """Return a final answer and every successful tool result used to produce it."""


class GeoAgentClient:
    """Small HTTP client for GeoAgent's public chat endpoint."""

    def __init__(
        self,
        *,
        chat_url: str,
        timeout: float,
        transport: httpx.AsyncBaseTransport | None = None,
    ) -> None:
        self.chat_url = chat_url
        self._client = httpx.AsyncClient(
            timeout=timeout,
            transport=transport,
            trust_env=False,
        )

    async def close(self) -> None:
        await self._client.aclose()

    async def answer(self, query: str, *, session_id: str) -> CollectedAnswer:
        # Cost safety: exactly one top-level GeoAgent request per workbook row.
        # Any HTTP/transport/payload failure is returned to the caller immediately.
        try:
            response = await self._client.post(
                self.chat_url,
                json={"session_id": session_id, "message": query},
            )
        except httpx.HTTPError as error:
            raise RuntimeError(f"Could not reach GeoAgent at {self.chat_url!r}: {error}") from error

        if response.is_error:
            raise RuntimeError(
                f"GeoAgent returned HTTP {response.status_code}: {response.text[:500]}"
            )
        try:
            payload = AgentResponse.model_validate(response.json())
        except (ValueError, ValidationError) as error:
            raise RuntimeError("GeoAgent returned an invalid response") from error
        answer = payload.answer.strip()
        if not answer:
            raise RuntimeError("GeoAgent returned an empty answer")
        successful_tool_results: list[SuccessfulToolResult] = []
        if payload.llm is not None:
            for step in payload.llm.trace.steps:
                if step.type.value != "observation" or not step.tool_name:
                    continue
                content = step.content.strip()
                if not content or content.startswith("ERROR["):
                    continue
                successful_tool_results.append(
                    SuccessfulToolResult(
                        tool_name=step.tool_name,
                        content=content,
                    )
                )
        return CollectedAnswer(
            answer=answer,
            successful_tool_results=tuple(successful_tool_results),
        )


def _cell_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _headers(sheet: Worksheet) -> dict[str, int]:
    headers: dict[str, int] = {}
    for column, cell in enumerate(sheet[1], start=1):
        name = _cell_text(cell.value)
        if not name:
            continue
        if name in headers:
            raise ValueError(f"Duplicate column {name!r} in worksheet {sheet.title!r}")
        headers[name] = column
    return headers


def _resolve_query_column(headers: dict[str, int], requested: str | None) -> str:
    if requested:
        if requested not in headers:
            raise ValueError(f"Query column {requested!r} was not found")
        return requested
    for candidate in QUERY_COLUMN_CANDIDATES:
        if candidate in headers:
            return candidate
    expected = ", ".join(QUERY_COLUMN_CANDIDATES)
    raise ValueError(f"Could not detect the query column; expected one of: {expected}")


def _ensure_answer_column(
    sheet: Worksheet,
    headers: dict[str, int],
    answer_column: str,
) -> int:
    existing = headers.get(answer_column)
    if existing is not None:
        return existing

    column = max(headers.values(), default=0) + 1
    target = sheet.cell(row=1, column=column, value=answer_column)
    if column > 1:
        source = sheet.cell(row=1, column=column - 1)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)
    return column


def _pending_rows(
    sheet: Worksheet,
    *,
    query_column: int,
    answer_column: int,
    answer_with_tools_column: int,
    overwrite: bool,
    start_row: int,
) -> tuple[list[PendingRow], int]:
    pending: list[PendingRow] = []
    skipped = 0
    for row in range(start_row, sheet.max_row + 1):
        query = _cell_text(sheet.cell(row=row, column=query_column).value)
        if not query:
            continue
        answer = _cell_text(sheet.cell(row=row, column=answer_column).value)
        answer_with_tools = _cell_text(sheet.cell(row=row, column=answer_with_tools_column).value)
        # Never spend tokens to repair a partially filled row implicitly. A user
        # must opt in with --overwrite, which makes the extra API call explicit.
        if (answer or answer_with_tools) and not overwrite:
            skipped += 1
            continue
        pending.append(PendingRow(excel_row=row, query=query))
    return pending, skipped


def _safe_session_part(value: str) -> str:
    cleaned = re.sub(r"[^a-zA-Z0-9_-]+", "-", value).strip("-")
    return cleaned[:40] or "sheet"


def _save_atomic(workbook: Workbook, path: Path) -> None:
    temporary = path.with_name(f".{path.stem}.{uuid.uuid4().hex}.tmp{path.suffix}")
    try:
        workbook.save(temporary)
        try:
            os.replace(temporary, path)
        except PermissionError as error:
            raise PermissionError(
                f"Cannot update {path.resolve()}. Close this workbook in Excel, "
                "LibreOffice, and any File Explorer preview pane, then run the "
                "command again. Existing answers will be skipped."
            ) from error
    finally:
        if temporary.exists():
            temporary.unlink()


async def fill_answers(
    *,
    workbook_path: Path,
    client: AnswerClient,
    sheet_name: str | None = None,
    query_column_name: str | None = None,
    answer_column_name: str = DEFAULT_ANSWER_COLUMN,
    answer_with_tools_column_name: str = DEFAULT_ANSWER_WITH_TOOLS_COLUMN,
    overwrite: bool = False,
    start_row: int = 2,
    limit: int | None = DEFAULT_API_CALL_LIMIT,
    dry_run: bool = False,
) -> FillResult:
    """Fill missing answers and checkpoint successful rows into ``workbook_path``."""
    if answer_column_name == answer_with_tools_column_name:
        raise ValueError("Answer columns must have different names")
    if start_row < 2:
        raise ValueError("start_row must be at least 2 because row 1 contains headers")

    keep_vba = workbook_path.suffix.lower() == ".xlsm"
    workbook = load_workbook(workbook_path, data_only=False, keep_vba=keep_vba)
    if sheet_name is not None:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Worksheet {sheet_name!r} was not found")
        sheet = workbook[sheet_name]
    else:
        sheet = workbook.active

    headers = _headers(sheet)
    query_name = _resolve_query_column(headers, query_column_name)
    answer_column = _ensure_answer_column(sheet, headers, answer_column_name)
    headers = _headers(sheet)
    answer_with_tools_column = _ensure_answer_column(
        sheet,
        headers,
        answer_with_tools_column_name,
    )
    eligible, skipped = _pending_rows(
        sheet,
        query_column=headers[query_name],
        answer_column=answer_column,
        answer_with_tools_column=answer_with_tools_column,
        overwrite=overwrite,
        start_row=start_row,
    )
    pending = eligible if limit is None else eligible[:limit]

    print(
        f"Workbook: {workbook_path}\n"
        f"Worksheet: {sheet.title}\n"
        f"Query column: {query_name}\n"
        f"Answer column: {answer_column_name}\n"
        f"Answer with tools column: {answer_with_tools_column_name}\n"
        f"Starting Excel row: {start_row}\n"
        f"Eligible rows: {len(eligible)}; selected API calls: {len(pending)}; "
        f"skipped filled/partial rows: {skipped}"
    )
    if dry_run:
        return FillResult(attempted=0, written=0, skipped=skipped, failed=0)

    run_id = uuid.uuid4().hex[:12]
    session_sheet = _safe_session_part(sheet.title)
    written = 0
    failed = 0
    attempted = 0
    for index, item in enumerate(pending, start=1):
        attempted += 1
        session_id = f"eval-{run_id}-{session_sheet}-row-{item.excel_row}"
        print(f"[{index}/{len(pending)}] row {item.excel_row}: sending...", flush=True)
        try:
            collected = await client.answer(item.query, session_id=session_id)
        except Exception as error:
            failed += 1
            print(f"[{index}/{len(pending)}] row {item.excel_row}: ERROR: {error}", flush=True)
            print("Stopping after the first failure to prevent further token spend.", flush=True)
            break
        sheet.cell(row=item.excel_row, column=answer_column, value=collected.answer)
        sheet.cell(
            row=item.excel_row,
            column=answer_with_tools_column,
            value=collected.with_tool_results(),
        )
        _save_atomic(workbook, workbook_path)
        written += 1
        print(f"[{index}/{len(pending)}] row {item.excel_row}: saved", flush=True)

    return FillResult(
        attempted=attempted,
        written=written,
        skipped=skipped,
        failed=failed,
    )


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Send Excel queries to GeoAgent and write final answers back to the same workbook."
        )
    )
    parser.add_argument("workbook", type=Path, help="Path to an .xlsx or .xlsm workbook")
    parser.add_argument("--sheet", help="Worksheet name; defaults to the active worksheet")
    parser.add_argument(
        "--query-column",
        help=(
            "Query column header. By default, auto-detects: " + ", ".join(QUERY_COLUMN_CANDIDATES)
        ),
    )
    parser.add_argument(
        "--answer-column",
        default=DEFAULT_ANSWER_COLUMN,
        help=f"Answer column header; created if absent (default: {DEFAULT_ANSWER_COLUMN})",
    )
    parser.add_argument(
        "--answer-with-tools-column",
        default=DEFAULT_ANSWER_WITH_TOOLS_COLUMN,
        help=(
            "Column for successful tool results followed by the final answer; "
            "created if absent "
            f"(default: {DEFAULT_ANSWER_WITH_TOOLS_COLUMN})"
        ),
    )
    parser.add_argument(
        "--agent-url",
        default=os.getenv("GEOAGENT_CHAT_URL", DEFAULT_AGENT_URL),
        help=f"GeoAgent chat endpoint (default: {DEFAULT_AGENT_URL})",
    )
    parser.add_argument("--timeout", type=float, default=600.0, help="Request timeout in seconds")
    parser.add_argument(
        "--limit",
        type=int,
        default=DEFAULT_API_CALL_LIMIT,
        help=(
            "Maximum number of API calls. Defaults to 1 for cost safety; "
            "bulk runs require an explicit value."
        ),
    )
    parser.add_argument(
        "--start-row",
        type=int,
        default=2,
        help="First Excel row to inspect, including the header row in numbering (default: 2)",
    )
    parser.add_argument(
        "--allow-bulk",
        action="store_true",
        help="Required together with --limit N when N is greater than 1",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Regenerate rows where either answer cell is already populated",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Validate the workbook and print the number of pending rows without modifying it",
    )
    return parser


def _validate_args(args: argparse.Namespace) -> None:
    if not args.workbook.is_file():
        raise SystemExit(f"Workbook does not exist: {args.workbook}")
    if args.workbook.suffix.lower() not in SUPPORTED_SUFFIXES:
        supported = ", ".join(sorted(SUPPORTED_SUFFIXES))
        raise SystemExit(f"Unsupported workbook format; expected one of: {supported}")
    if args.timeout <= 0:
        raise SystemExit("--timeout must be greater than zero")
    if args.limit is not None and args.limit < 1:
        raise SystemExit("--limit must be at least 1")
    if args.start_row < 2:
        raise SystemExit("--start-row must be at least 2 because row 1 contains headers")
    if args.limit is not None and args.limit > 1 and not args.allow_bulk:
        raise SystemExit(
            "Bulk API calls are disabled by default. Run --dry-run first, then pass "
            "--limit N --allow-bulk explicitly."
        )
    if not args.answer_column.strip():
        raise SystemExit("--answer-column cannot be empty")
    if not args.answer_with_tools_column.strip():
        raise SystemExit("--answer-with-tools-column cannot be empty")
    if args.answer_column == args.answer_with_tools_column:
        raise SystemExit("--answer-column and --answer-with-tools-column must be different")


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    _validate_args(args)
    client = GeoAgentClient(
        chat_url=args.agent_url,
        timeout=args.timeout,
    )

    async def execute() -> FillResult:
        try:
            return await fill_answers(
                workbook_path=args.workbook,
                client=client,
                sheet_name=args.sheet,
                query_column_name=args.query_column,
                answer_column_name=args.answer_column,
                answer_with_tools_column_name=args.answer_with_tools_column,
                overwrite=args.overwrite,
                start_row=args.start_row,
                limit=args.limit,
                dry_run=args.dry_run,
            )
        finally:
            await client.close()

    try:
        result = asyncio.run(execute())
    except (OSError, ValueError) as error:
        raise SystemExit(str(error)) from error

    print(
        f"Done: attempted={result.attempted}, written={result.written}, "
        f"skipped={result.skipped}, failed={result.failed}"
    )
    return 1 if result.failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
